#!/usr/bin/env python3
"""
Generate a Word document: Algorithm Design Chapter for
Space Station Robotic Arm Task Planning Based on EUROPA.
"""

from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
import os

def set_cell_shading(cell, color):
    shading = cell._element.get_or_add_tcPr()
    shading_elem = shading.makeelement(qn('w:shd'), {
        qn('w:val'): 'clear',
        qn('w:color'): 'auto',
        qn('w:fill'): color
    })
    shading.append(shading_elem)

def add_code_block(doc, code_text):
    for line in code_text.strip().split('\n'):
        p = doc.add_paragraph()
        p.style = doc.styles['code']
        p.add_run(line)

def create_document():
    doc = Document()

    # ── Style setup ──
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(6)

    # Chinese font fallback
    rFonts = style.element.rPr.rFonts if style.element.rPr is not None else None
    if rFonts is None:
        from docx.oxml import OxmlElement
        rPr = style.element.get_or_add_rPr()
        rFonts = OxmlElement('w:rFonts')
        rFonts.set(qn('w:eastAsia'), '宋体')
        rPr.append(rFonts)

    for level in range(1, 4):
        hs = doc.styles[f'Heading {level}']
        hs.font.color.rgb = RGBColor(0, 0, 0)
        hs.font.bold = True
        if level == 1:
            hs.font.size = Pt(16)
        elif level == 2:
            hs.font.size = Pt(14)
        else:
            hs.font.size = Pt(12)

    # Code style
    code_style = doc.styles.add_style('code', WD_STYLE_TYPE.PARAGRAPH)
    code_style.font.name = 'Consolas'
    code_style.font.size = Pt(9)
    code_style.paragraph_format.space_before = Pt(0)
    code_style.paragraph_format.space_after = Pt(0)
    code_style.paragraph_format.line_spacing = 1.15
    code_style.paragraph_format.left_indent = Cm(1)

    # ════════════════════════════════════════════
    # TITLE
    # ════════════════════════════════════════════
    title = doc.add_heading('第X章 算法设计', level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ════════════════════════════════════════════
    # 3.1 Problem formulation
    # ════════════════════════════════════════════
    doc.add_heading('X.1 问题形式化描述', level=2)

    doc.add_paragraph(
        '空间站机械臂移动任务规划问题可形式化为一个带时序约束的规划问题。'
        '机械臂需要沿预设轨道从初始位置移动到目标位置，移动过程受到通信窗口和遮挡约束的限制。'
    )

    doc.add_heading('X.1.1 问题定义', level=3)

    doc.add_paragraph('定义规划问题 P = ⟨S, A, T, C, s₀, g⟩，其中：')
    items = [
        'S = {base1, on_900_base2, base2, on_900_base3, base3} 为位置状态空间',
        'A = {MoveToOn900Base2, MoveToBase2, MoveToOn900Base3, MoveToBase3} 为动作集合',
        'T = {CommWindow, OcclusionBase2} 为时序约束时间线集合',
        'C 为约束集合（通信窗口约束、遮挡约束、时序约束）',
        's₀ = AtBase1 为初始状态',
        'g = AtBase3 为目标状态',
    ]
    for item in items:
        doc.add_paragraph(item, style='List Bullet')

    doc.add_heading('X.1.2 动作持续时间模型', level=3)

    doc.add_paragraph(
        '每个移动动作的持续时间由梯形速度曲线决定。'
        '给定最大速度 v_max 和最大加速度 a_max，角度距离为 d 的移动段持续时间为：'
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('T(d) = 2·(v_max / a_max) + (d − v_max² / a_max) / v_max')
    run.italic = True

    doc.add_paragraph(
        '其中 v_max = 0.5°/s，a_max = 0.2°/s²。该公式包含加速段、匀速段和减速段三个阶段。'
    )

    # Duration table
    table = doc.add_table(rows=5, cols=4, style='Table Grid')
    headers = ['移动段', '角度距离 (°)', '持续时间 (s)', '公式展开']
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
        set_cell_shading(cell, 'D9E2F3')
    data = [
        ['base1 → on_900_base2', '260', '523', '5 + (260−1.25)/0.5'],
        ['on_900_base2 → base2', '100', '203', '5 + (100−1.25)/0.5'],
        ['base2 → on_900_base3', '290', '583', '5 + (290−1.25)/0.5'],
        ['on_900_base3 → base3', '100', '203', '5 + (100−1.25)/0.5'],
    ]
    for ri, row_data in enumerate(data):
        for ci, val in enumerate(row_data):
            table.rows[ri+1].cells[ci].text = val

    doc.add_paragraph()

    doc.add_heading('X.1.3 约束建模', level=3)

    doc.add_paragraph('系统包含以下类型的约束：')
    constraints = [
        ('通信窗口约束', '每个移动动作必须完全包含在某个通信可用窗口内。形式化为 contained_by(Move, InComms)，即 InComms.start ≤ Move.start ∧ Move.end ≤ InComms.end。'),
        ('遮挡约束', '从 base2 出发的移动动作必须在遮挡不活跃期间执行。形式化为 contained_by(MoveFromBase2, OcclusionInactive)。'),
        ('因果约束', '每个移动动作要求前置位置状态（met_by）和产生后置位置状态（meets），构成因果链。'),
        ('持续时间约束', '每个移动动作的 start + duration = end，duration 为预计算的固定值。'),
    ]
    for name, desc in constraints:
        p = doc.add_paragraph()
        run = p.add_run(f'{name}：')
        run.bold = True
        p.add_run(desc)

    # ════════════════════════════════════════════
    # 3.2 Algorithm overview
    # ════════════════════════════════════════════
    doc.add_heading('X.2 核心算法：缺陷导向时序回溯搜索', level=2)

    doc.add_paragraph(
        '本文采用 EUROPA 框架的缺陷导向时序回溯搜索算法（Flaw-Directed Temporal Backtracking Search）求解上述规划问题。'
        '该算法不同于传统的前向状态空间搜索，而是从一个包含初始状态和目标的不完整部分计划出发，'
        '通过反复发现并修复计划中的"缺陷"来逐步精化计划，直到所有缺陷消除，得到完整可执行的计划。'
    )

    doc.add_heading('X.2.1 部分计划与缺陷', level=3)

    doc.add_paragraph(
        '部分计划（Partial Plan）是一个四元组 PP = ⟨Tokens, Constraints, Timeline, Flaws⟩。'
        'Token 表示时间轴上的动作或状态实例，约束限制 Token 之间的时序和参数关系，'
        '时间线（Timeline）强制同一资源上的 Token 不重叠。Flaws 是计划中需要修复的缺陷集合。'
    )

    doc.add_paragraph('算法识别三类缺陷：')

    table2 = doc.add_table(rows=4, cols=3, style='Table Grid')
    t2_headers = ['缺陷类型', '定义', '解决策略']
    for i, h in enumerate(t2_headers):
        cell = table2.rows[0].cells[i]
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
        set_cell_shading(cell, 'D9E2F3')
    t2_data = [
        ['开放条件\n(Open Condition)', '存在状态为 INACTIVE 的 Token\n（未被放置到时间线上）', '① 激活(Activate)：放置到时间线\n② 合并(Merge)：与已有 Token 统一\n③ 拒绝(Reject)：排除此可能性'],
        ['时序威胁\n(Threat)', '时间线上有 Token 的前后\n顺序未确定', '选择一个合法的排序位置\n(predecessor, successor) 对'],
        ['未绑定变量\n(Unbound Variable)', '已激活 Token 的参数变量\n域未收窄为单值', '从变量域中选取一个值\n进行绑定(Specify)'],
    ]
    for ri, row_data in enumerate(t2_data):
        for ci, val in enumerate(row_data):
            table2.rows[ri+1].cells[ci].text = val

    doc.add_paragraph()

    doc.add_heading('X.2.2 规则引擎与因果推理', level=3)

    doc.add_paragraph(
        '当一个动作 Token（如 MoveToBase2）被激活时，规则引擎自动触发该动作的规则体，'
        '创建子 Token（前置条件和后置效果）并添加时序约束。'
        '这些新创建的子 Token 以 INACTIVE 状态进入计划，成为新的开放条件缺陷，驱动求解器继续搜索。'
        '这一机制实现了目标驱动的反向链式推理（Backward Chaining）。'
    )

    doc.add_paragraph('以 MoveToBase2 为例，规则触发产生：')
    rule_items = [
        'met_by(AtOn900Base2) → 创建前置条件 Token',
        'meets(AtBase2) → 创建后置效果 Token',
        'contained_by(InComms) → 创建通信约束 Token',
        'eq(duration, 203) → 添加持续时间约束',
    ]
    for item in rule_items:
        doc.add_paragraph(item, style='List Bullet')

    # ════════════════════════════════════════════
    # 3.3 Main algorithm pseudocode
    # ════════════════════════════════════════════
    doc.add_heading('X.3 算法伪代码', level=2)

    doc.add_heading('X.3.1 主求解循环', level=3)

    doc.add_paragraph(
        '算法1给出了主求解循环的伪代码。'
        '求解器维护一个决策栈（Decision Stack）记录已做出的决策，支持时序回溯。'
        '每一步先进行约束传播检查一致性，然后选择最高优先级的缺陷，'
        '做出决策并再次传播。如果传播后不一致，则触发回溯。'
    )

    p = doc.add_paragraph()
    run = p.add_run('算法1：缺陷导向时序回溯搜索（主循环）')
    run.bold = True

    code1 = """
输入: 部分计划 PP = ⟨Tokens, Constraints, Timeline, Flaws⟩
      最大步数 maxSteps, 最大深度 maxDepth
输出: 完整计划或"无解"

函数 SOLVE(PP, maxSteps, maxDepth):
  DecisionStack ← ∅
  stepCount ← 0
  WHILE stepCount < maxSteps:
    // 步骤1: 约束传播
    consistent ← PROPAGATE(PP.ConstraintEngine)
    IF ¬consistent:
      RETURN "无解"    // 初始状态已矛盾

    // 步骤2: 选择缺陷并创建决策点
    decision ← ALLOCATE_DECISION(PP.FlawManagers)
    IF decision = ∅:
      RETURN PP        // 无缺陷 → 求解成功

    IF stepCount ≥ maxSteps OR |DecisionStack| ≥ maxDepth:
      RETURN "超时"

    // 步骤3: 执行决策
    decision.EXECUTE()
    PROPAGATE(PP.ConstraintEngine)
    stepCount ← stepCount + 1

    // 步骤4: 检查一致性
    IF CONSISTENT(PP):
      DecisionStack.PUSH(decision)
      CONTINUE         // 进入下一轮
    ELSE:
      // 步骤5: 回溯
      exhausted ← BACKTRACK(decision, DecisionStack)
      IF exhausted:
        RETURN "无解"  // 搜索空间耗尽

  RETURN "超时"
"""
    add_code_block(doc, code1)
    doc.add_paragraph()

    doc.add_heading('X.3.2 缺陷选择算法', level=3)

    doc.add_paragraph(
        '算法2给出了缺陷选择的伪代码。'
        '算法首先检查是否存在零承诺决策（仅有唯一选择的缺陷），'
        '然后在所有缺陷管理器中选择优先级最高的缺陷。'
    )

    p = doc.add_paragraph()
    run = p.add_run('算法2：缺陷选择（ALLOCATE_DECISION）')
    run.bold = True

    code2 = """
输入: 缺陷管理器集合 FM = {ThreatMgr, OpenCondMgr, UnboundVarMgr}
输出: 决策点 decision 或 ∅

函数 ALLOCATE_DECISION(FM):
  // 优先处理零承诺决策（仅一个选择的缺陷）
  FOR EACH fm ∈ FM:
    d ← fm.NEXT_ZERO_COMMITMENT()
    IF d ≠ ∅:
      d.INITIALIZE()
      RETURN d

  // 在所有管理器中选择最高优先级缺陷
  bestDecision ← ∅
  bestPriority ← +∞
  FOR EACH fm ∈ FM:
    FOR EACH flaw ∈ fm.ITERATOR():
      IF fm.FILTERED(flaw):
        CONTINUE
      priority ← fm.GET_PRIORITY(flaw)
      IF priority < bestPriority:
        bestDecision ← fm.CREATE_DECISION(flaw)
        bestPriority ← priority

  IF bestDecision ≠ ∅:
    bestDecision.INITIALIZE()
  RETURN bestDecision
"""
    add_code_block(doc, code2)
    doc.add_paragraph()

    doc.add_heading('X.3.3 回溯算法', level=3)

    doc.add_paragraph(
        '算法3给出了时序回溯的伪代码。'
        '回溯采用时序回溯策略（Chronological Backtracking），'
        '从决策栈顶开始逐个撤销决策，直到找到还有可选分支的决策点。'
    )

    p = doc.add_paragraph()
    run = p.add_run('算法3：时序回溯（BACKTRACK）')
    run.bold = True

    code3 = """
输入: 当前决策 activeDec, 决策栈 Stack
输出: exhausted ∈ {true, false}

函数 BACKTRACK(activeDec, Stack):
  backtracking ← true
  WHILE backtracking ∧ (activeDec ≠ ∅ ∨ Stack ≠ ∅):

    IF activeDec = ∅ ∧ Stack ≠ ∅:
      activeDec ← Stack.POP()

    // 撤销已执行的决策
    IF activeDec.IS_EXECUTED():
      activeDec.UNDO()
      // 约束网络自动回退到决策前的状态

    // 检查是否还有其他选择
    IF activeDec.HAS_NEXT_CHOICE():
      backtracking ← false    // 找到可继续的分支
    ELSE:
      DELETE activeDec
      activeDec ← ∅           // 继续向上回溯

  RETURN backtracking          // true 表示搜索空间耗尽
"""
    add_code_block(doc, code3)
    doc.add_paragraph()

    # ════════════════════════════════════════════
    # 3.4 STN Algorithm
    # ════════════════════════════════════════════
    doc.add_heading('X.4 简单时序网络与增量 Dijkstra 传播', level=2)

    doc.add_paragraph(
        '算法的核心效率来源于简单时序网络（Simple Temporal Network, STN）的增量传播。'
        'STN 维护所有时间变量之间的距离约束，通过增量 Dijkstra 算法高效地传播约束变化，'
        '避免了每次决策后的全量重计算。'
    )

    doc.add_heading('X.4.1 约束编码', level=3)

    doc.add_paragraph(
        '时序约束 lb ≤ B − A ≤ ub 被编码为有向距离图中的两条边：'
    )
    enc_items = [
        'A → B，权重 ub（上界边）：表示 B − A ≤ ub',
        'B → A，权重 −lb（下界边）：表示 A − B ≤ −lb，即 B − A ≥ lb',
    ]
    for item in enc_items:
        doc.add_paragraph(item, style='List Bullet')

    doc.add_paragraph(
        '时间点 T 相对于原点 O 的界通过最短路径计算：'
        'upperBound(T) = 从 O 到 T 的最短路径长度；'
        'lowerBound(T) = −(从 T 到 O 的最短路径长度)。'
        '若存在负环，则系统不一致（矛盾）。'
    )

    doc.add_heading('X.4.2 Johnson 重标号与增量 Dijkstra', level=3)

    doc.add_paragraph(
        '由于下界边的权重为负值，标准 Dijkstra 算法不可直接使用。'
        '算法采用 Johnson 重标号技术：首先通过 Bellman-Ford 算法计算每个节点的势函数 π(v)，'
        '然后将边权重标号为 w\'(u,v) = w(u,v) + π(u) − π(v) ≥ 0，'
        '使得所有边权非负，从而可以使用 Dijkstra 算法进行高效传播。'
    )

    doc.add_paragraph(
        '增量传播的关键优化在于：添加一条新约束时，'
        '算法首先通过 startNode() 函数检测新边是否能改进任一端点的距离。'
        '若不能改进，则无需传播（O(1)跳过）；若能改进，则仅从受影响的节点开始局部 Dijkstra 传播。'
    )

    p = doc.add_paragraph()
    run = p.add_run('算法4：增量传播（INC_PROPAGATE）')
    run.bold = True

    code4 = """
输入: 新约束的两个端点 src, targ
效果: 更新所有受影响时间点的上下界

函数 INC_PROPAGATE(src, targ):
  IF 有待处理的删除 ∨ 已不一致:
    RETURN    // 需要完全重传播

  // 阶段1: 一致性检查（增量 Bellman-Ford）
  startPt ← START_NODE(src.potential, targ.potential, forward)
  IF startPt ≠ ∅:
    consistent ← INC_BELLMAN_FORD(startPt)
    IF ¬consistent:
      MARK_INCONSISTENT()
      RETURN

  // 阶段2: 上界传播（增量 Dijkstra 前向）
  startPt ← START_NODE(src.upperBound, targ.upperBound, forward)
  IF startPt ≠ ∅:
    Queue.INSERT(startPt)
    INC_DIJKSTRA_FORWARD(Queue)

  // 阶段3: 下界传播（增量 Dijkstra 后向）
  startPt ← START_NODE(−src.lowerBound, −targ.lowerBound, backward)
  IF startPt ≠ ∅:
    Queue.INSERT(startPt)
    INC_DIJKSTRA_BACKWARD(Queue)
"""
    add_code_block(doc, code4)
    doc.add_paragraph()

    p = doc.add_paragraph()
    run = p.add_run('算法5：增量 Dijkstra 前向传播（上界）')
    run.bold = True

    code5 = """
输入: 优先队列 Queue（含初始受影响节点）
效果: 更新所有受影响节点的 upperBound

函数 INC_DIJKSTRA_FORWARD(Queue):
  WHILE Queue ≠ ∅:
    node ← Queue.POP_MIN()       // 弹出键值最小的节点
    FOR EACH edge ∈ node.outEdges:
      next ← edge.to
      newDist ← node.upperBound + edge.weight
      IF newDist < next.upperBound:
        // 发现更紧的上界
        next.upperBound ← newDist
        next.depth ← node.depth + 1
        IF next.depth > |V|:       // 环路检测
          MARK_INCONSISTENT()
          RETURN
        // Johnson 重标号确保键值非负
        key ← newDist − next.potential
        Queue.INSERT(next, key)
        MARK_UPDATED(next)
"""
    add_code_block(doc, code5)
    doc.add_paragraph()

    p = doc.add_paragraph()
    run = p.add_run('算法6：增量 Dijkstra 后向传播（下界）')
    run.bold = True

    code6 = """
输入: 优先队列 Queue（含初始受影响节点）
效果: 更新所有受影响节点的 lowerBound

函数 INC_DIJKSTRA_BACKWARD(Queue):
  WHILE Queue ≠ ∅:
    node ← Queue.POP_MIN()
    FOR EACH edge ∈ node.inEdges:    // 注意: 遍历入边
      next ← edge.from               // 注意: 反向传播
      newDist ← (−node.lowerBound) + edge.weight
      IF newDist < (−next.lowerBound):
        next.lowerBound ← −newDist
        next.depth ← node.depth + 1
        IF next.depth > |V|:
          MARK_INCONSISTENT()
          RETURN
        key ← newDist + next.potential
        Queue.INSERT(next, key)
        MARK_UPDATED(next)
"""
    add_code_block(doc, code6)
    doc.add_paragraph()

    doc.add_heading('X.4.3 约束引擎与STN的双向同步', level=3)

    doc.add_paragraph(
        '时序传播器（Temporal Propagator）作为约束引擎（CE）与 STN 之间的桥梁，'
        '负责双向同步：CE 中变量域的变化通过 updateTnet() 传递到 STN，'
        'STN 传播后的新界通过 updateCnet() 推回 CE 中收窄变量域。'
        '这一机制确保时序推理与非时序约束传播协同工作。'
    )

    # ════════════════════════════════════════════
    # 3.5 Complexity
    # ════════════════════════════════════════════
    doc.add_heading('X.5 算法复杂度分析', level=2)

    table3 = doc.add_table(rows=6, cols=3, style='Table Grid')
    t3h = ['操作', '时间复杂度', '说明']
    for i, h in enumerate(t3h):
        cell = table3.rows[0].cells[i]
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
        set_cell_shading(cell, 'D9E2F3')
    t3d = [
        ['添加时序约束', 'O((V+E) log V)', '增量 Dijkstra，V为时间点数，E为边数'],
        ['删除时序约束', 'O(VE)', '需完全重传播（Bellman-Ford + Dijkstra）'],
        ['一致性检查', 'O(1)', '传播时已检测，直接读取标志'],
        ['查询时间点界', 'O(1)', '直接读取 upperBound/lowerBound'],
        ['完整求解', 'O(S · (V+E) log V)', 'S为搜索步数，每步一次增量传播'],
    ]
    for ri, row_data in enumerate(t3d):
        for ci, val in enumerate(row_data):
            table3.rows[ri+1].cells[ci].text = val

    doc.add_paragraph()

    doc.add_paragraph(
        '对于本文的机械臂移动规划问题，时间点数 V ≈ 40，边数 E ≈ 80，'
        '求解器在 45 步内完成求解。每次增量传播在微秒级完成，整体求解时间约 75 毫秒。'
    )

    # ════════════════════════════════════════════
    # 3.6 Walkthrough
    # ════════════════════════════════════════════
    doc.add_heading('X.6 算法执行过程示例', level=2)

    doc.add_paragraph(
        '以下展示算法求解机械臂从 base1 移动到 base3 的关键决策步骤：'
    )

    table4 = doc.add_table(rows=10, cols=4, style='Table Grid')
    t4h = ['步骤', '缺陷类型', '决策', '效果']
    for i, h in enumerate(t4h):
        cell = table4.rows[0].cells[i]
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
        set_cell_shading(cell, 'D9E2F3')
    t4d = [
        ['1', 'OpenCondition', 'ACTIVATE AtBase3', '目标放置到Arm时间线'],
        ['3-4', 'OpenCondition', 'ACTIVATE MoveToBase3', '触发规则，创建子Token'],
        ['5-6', 'OpenCondition', 'MERGE InComms→comm3', '约束Move在[1500,2900]内'],
        ['13-14', 'OpenCondition', 'ACTIVATE MoveToOn900Base3', '创建遮挡约束子Token'],
        ['17-18', 'OpenCondition', 'MERGE Inactive→occ3', '约束Move在[1650,3000]内'],
        ['23-24', 'OpenCondition', 'MERGE InComms→comm1(回溯)', '不一致→回溯→改为comm2'],
        ['25-29', '回溯+重试', 'MERGE InComms→comm2', '约束MoveToBase2在[690,1200]'],
        ['35-36', 'OpenCondition', 'MERGE InComms→comm1', '约束Move1在[0,660]'],
        ['37-44', 'Threat', '时间线排序', '确定Arm上所有Token顺序'],
    ]
    for ri, row_data in enumerate(t4d):
        for ci, val in enumerate(row_data):
            table4.rows[ri+1].cells[ci].text = val

    doc.add_paragraph()

    doc.add_paragraph(
        '经过45步搜索（含回溯），求解器生成的最终计划为：'
    )

    table5 = doc.add_table(rows=10, cols=4, style='Table Grid')
    t5h = ['序号', 'Token', '时间区间', '说明']
    for i, h in enumerate(t5h):
        cell = table5.rows[0].cells[i]
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
        set_cell_shading(cell, 'D9E2F3')
    t5d = [
        ['1', 'AtBase1', '[0, 137]', '等待通信窗口'],
        ['2', 'MoveToOn900Base2', '[137, 660]', '523s, 在comm1[0,660]内'],
        ['3', 'AtOn900Base2', '[660, 690]', '等待通信恢复'],
        ['4', 'MoveToBase2', '[690, 893]', '203s, 在comm2[690,1200]内'],
        ['5', 'AtBase2', '[893, 1650]', '等待遮挡结束+通信恢复'],
        ['6', 'MoveToOn900Base3', '[1650, 2233]', '583s, 遮挡不活跃+comm3'],
        ['7', 'AtOn900Base3', '[2233, 2436]', '短暂等待'],
        ['8', 'MoveToBase3', '[2436, 2639]', '203s, 在comm3[1500,2900]内'],
        ['9', 'AtBase3', '[2639, 3000]', '到达目标'],
    ]
    for ri, row_data in enumerate(t5d):
        for ci, val in enumerate(row_data):
            table5.rows[ri+1].cells[ci].text = val

    doc.add_paragraph()

    # ════════════════════════════════════════════
    # 3.7 Summary
    # ════════════════════════════════════════════
    doc.add_heading('X.7 本章小结', level=2)

    doc.add_paragraph(
        '本章提出了基于 EUROPA 框架的空间站机械臂任务规划算法。'
        '该算法的核心特点包括：'
    )
    summary_items = [
        '缺陷导向搜索：从不完整部分计划出发，通过修复缺陷逐步精化，避免了前向搜索的状态爆炸问题。',
        '规则驱动的因果推理：动作激活时自动创建前置条件和后置效果，实现目标驱动的反向链式推理。',
        '增量时序传播：基于 Johnson 重标号的增量 Dijkstra 算法，单次约束添加的传播复杂度为 O((V+E) log V)，避免了全量重计算。',
        '约束传播剪枝：每次决策后立即传播约束，尽早发现不一致，大幅减少无效搜索分支。',
        '时间线排序机制：Timeline 类强制 Token 不重叠，自动生成合法的排序选择，确保计划在时序上的可行性。',
    ]
    for item in summary_items:
        doc.add_paragraph(item, style='List Bullet')

    doc.add_paragraph(
        '实验结果表明，该算法能在 45 步、75 毫秒内完成包含 4 段移动、'
        '3 个通信窗口和 2 个遮挡区间的机械臂规划问题，验证了算法的有效性和效率。'
    )

    # ── Save ──
    out_path = '/opt/cursor/artifacts/algorithm_design_chapter.docx'
    doc.save(out_path)
    local_path = '/workspace/examples/JxbMove/algorithm_design_chapter.docx'
    doc.save(local_path)
    print(f"Saved: {out_path}")
    print(f"Saved: {local_path}")

if __name__ == '__main__':
    create_document()
