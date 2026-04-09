#!/usr/bin/env python3
"""
Generate Chapter 2 v2: No pseudocode blocks.
Algorithms described in step-by-step academic narrative style.
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn

OUT = '/workspace/examples/experiment_final'

def shd(c, cl):
    pr = c._element.get_or_add_tcPr()
    pr.append(pr.makeelement(qn('w:shd'),{qn('w:val'):'clear',qn('w:color'):'auto',qn('w:fill'):cl}))

def htb(doc, hd, rows):
    t = doc.add_table(rows=1+len(rows), cols=len(hd), style='Table Grid')
    for i, h in enumerate(hd):
        c = t.rows[0].cells[i]; c.text = h; c.paragraphs[0].runs[0].bold = True; shd(c, 'D9E2F3')
    for ri, row in enumerate(rows):
        for ci, v in enumerate(row): t.rows[ri+1].cells[ci].text = str(v)

def fm(doc, text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text); r.italic = True; return p

def br(p, text):
    r = p.add_run(text); r.bold = True; return r

def create():
    doc = Document()
    s = doc.styles['Normal']; s.font.name = 'Times New Roman'; s.font.size = Pt(12)
    s.paragraph_format.line_spacing = 1.5; s.paragraph_format.space_after = Pt(6)
    rPr = s.element.get_or_add_rPr()
    rPr.append(rPr.makeelement(qn('w:rFonts'), {qn('w:eastAsia'): '宋体'}))
    for lv in range(1, 4):
        h = doc.styles[f'Heading {lv}']; h.font.color.rgb = RGBColor(0,0,0)
        h.font.bold = True; h.font.size = Pt([0, 16, 14, 12][lv])

    # ════════════════════════════════════════
    title = doc.add_heading('第2章 空间站机械臂作业任务规划建模与求解', level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(
        '空间站机械臂在轨作业任务规划问题涉及离散的任务状态转换与连续的时序约束推理，'
        '是一类典型的混合约束满足问题。本章首先建立问题的总体描述（2.1节），'
        '然后给出时间线规划系统的形式化定义（2.2节），在此基础上对约束体系进行系统建模（2.3节），'
        '利用NDDL语言完成编码实现（2.4节），设计动作规则体系（2.5节），'
        '最后给出基于缺陷导向搜索与简单时序网络传播的求解算法（2.6节）。')

    # ═══════════════════
    # 2.1
    # ═══════════════════
    doc.add_heading('2.1 问题概述', level=2)

    doc.add_paragraph(
        '空间站机械臂在轨作业任务规划问题可抽象为一类带时序约束和资源约束的调度问题。'
        '机械臂需执行一系列预定义的操作子任务，每个子任务具有固定的持续时间和严格的执行顺序。'
        '任务执行过程受到两类环境约束的限制：其一，所有操作必须在中继卫星提供的通信可见窗口内完成；'
        '其二，当机械臂在SWA（转位机构）附近作业时，臂杆可能遮挡空间站天线与中继卫星之间的通信链路，'
        '须回避遮挡区间。')

    doc.add_paragraph(
        '上述问题的核心挑战在于：（1）通信窗口的周期性中断（每轨道圈约42分钟不可用）'
        '限制了任务的连续执行；（2）遮挡约束与通信约束的空间耦合——遮挡仅在通信可用期间有实际影响；'
        '（3）提前量约束（大臂重启需在运动前12小时、小臂重启需在运动前5小时完成）'
        '跨越多个轨道圈次，增加了调度的时间跨度；'
        '（4）部分阶段（如平台设置，总时长3530秒）超过单个通信窗口容量（约3400秒），'
        '需自动拆分到相邻窗口执行。规划算法需在满足上述所有约束的前提下确定各子任务的可行执行时间。')

    # ═══════════════════
    # 2.2
    # ═══════════════════
    doc.add_heading('2.2 时间线规划系统形式化定义', level=2)

    doc.add_paragraph(
        '本节采用EUROPA框架的时间线规划范式对问题进行形式化建模。'
        '时间线规划的核心思想是：将系统中每个资源或对象建模为一条时间线，'
        '时间线上的Token（令牌）表示该资源在某时段内的状态或动作，'
        'Token形成有序不重叠序列，自然表达了"同一资源在同一时刻只能处于一种状态"的语义。')

    # Def 1
    doc.add_heading('2.2.1 时间线规划系统', level=3)
    p = doc.add_paragraph(); br(p, '定义1（时间线规划系统）')
    doc.add_paragraph('时间线规划系统（Timeline Planning System, TPS）𝒮 是一个五元组：')
    fm(doc, '𝒮 = (𝒪, ℒ, 𝒯, 𝒞, ℛ)                (1)')
    doc.add_paragraph('式中，各元素定义如下：')
    for sym, desc in [
        ('𝒪 = {o₁, o₂, …, oₘ}',
         '为有限对象集合。每个对象oᵢ通过双射L:𝒪→ℒ关联一条时间线L(oᵢ)=Lᵢ。'),
        ('ℒ = {L₁, L₂, …, Lₘ}',
         '为时间线集合。每条时间线上的Token形成有序不重叠序列，'
         '即对同一时间线上任意相邻Token τⱼ和τⱼ₊₁，有end(τⱼ)≤start(τⱼ₊₁)。'),
        ('𝒯', '为Token类型集合，分为谓词（predicate）和动作（action）两类。'
         '每个Token τ具有类型type(τ)、起始时间start(τ)、结束时间end(τ)和持续时间dur(τ)=end(τ)−start(τ)。'),
        ('𝒞', '为约束集合，包含时序约束和参数约束，详见2.3节。'),
        ('ℛ', '为规则集合，定义动作Token被激活时自动创建的子Token和约束，详见2.5节。'),
    ]:
        p = doc.add_paragraph(); br(p, f'{sym} '); p.add_run(desc)

    # Def 2
    doc.add_heading('2.2.2 部分计划与缺陷', level=3)
    p = doc.add_paragraph(); br(p, '定义2（部分计划）')
    doc.add_paragraph('部分计划PP是一个三元组PP=(Tokens, Active, Constraints)，'
        '其中Tokens为当前计划中所有Token的集合，Active⊆Tokens为已放置到时间线上的Token子集，'
        'Constraints为当前所有约束的集合。当部分计划中不存在缺陷时，PP为完整计划。')

    p = doc.add_paragraph(); br(p, '定义3（缺陷）')
    doc.add_paragraph('部分计划中的缺陷分为三类：'
        '（1）开放条件——Token存在于计划中但尚未放置到时间线上；'
        '（2）时序威胁——已激活Token在其时间线上与其他Token的排序未确定；'
        '（3）未绑定变量——已激活Token的某参数变量域尚未收窄为单值。'
        '求解算法通过逐一消除缺陷来精化部分计划，直至得到完整计划。')

    # Def 4
    doc.add_heading('2.2.3 简单时序网络', level=3)
    p = doc.add_paragraph(); br(p, '定义4（简单时序网络）')
    doc.add_paragraph('简单时序网络（STN）N=(V, E, O)是一个有向加权距离图，'
        '其中V为时间点集合（每个Token贡献start、end、duration三个时间点），'
        'E为有向加权边集合，O为参考原点。')
    doc.add_paragraph('时序约束lb≤Tⱼ−Tᵢ≤ub在STN中编码为两条有向边：')
    fm(doc, 'Tᵢ →^{ub} Tⱼ（上界边）和 Tⱼ →^{−lb} Tᵢ（下界边）                (2)')

    p = doc.add_paragraph(); br(p, '性质1 ')
    p.add_run('时间点T的上界等于从原点O到T的最短路径长度（最晚可能时刻），'
        '下界等于从T到O的最短路径长度的相反数（最早可能时刻）。'
        '若距离图中存在负权环路，则约束系统不一致。')

    # Def 5
    doc.add_heading('2.2.4 机械臂作业规划问题', level=3)
    p = doc.add_paragraph(); br(p, '定义5（机械臂作业规划问题）')
    doc.add_paragraph('机械臂作业规划问题是TPS上的约束满足问题𝒫=(𝒮, s₀, 𝒢, H)，'
        '其中𝒮为时间线规划系统（定义1），对象集合𝒪={Arm, CommWindow, OcclusionSWA}；'
        's₀为初始部分计划，包含环境时间线上的Fact Token和Phase 1子任务的固定放置；'
        '𝒢为目标条件；H为规划范围上界。')
    doc.add_paragraph('问题的解是一个完整计划PP*，满足：（1）无缺陷；'
        '（2）所有已激活Token时序一致；'
        '（3）所有操作Token的时间区间被包含于某个通信可用区间内。')

    # ═══════════════════
    # 2.3
    # ═══════════════════
    doc.add_heading('2.3 约束建模', level=2)

    doc.add_heading('2.3.1 顺序约束', level=3)
    doc.add_paragraph('同一阶段内相邻子任务满足紧邻约束（Allen关系中的meets）：')
    fm(doc, 'C_SEQ = { (τᵢ, τⱼ) | end(τᵢ) = start(τⱼ) }                (3)')
    doc.add_paragraph('阶段间的衔接允许等待间隔，满足弱顺序约束end(τᵢ) ≤ start(τⱼ)。')

    doc.add_heading('2.3.2 通信窗口约束', level=3)
    doc.add_paragraph('通信窗口序列由SGP4轨道传播模型计算确定。'
        '通信窗口约束要求每个操作子任务完全包含在某个通信可见窗口内：')
    fm(doc, '∀τᵢ ∈ 𝒯_op, ∃k: aₖ ≤ start(τᵢ) ∧ end(τᵢ) ≤ bₖ                (4)')

    doc.add_heading('2.3.3 遮挡约束', level=3)
    doc.add_paragraph('遮挡状态通过SGP4轨道传播联合DH运动学正解计算确定。'
        'SWA附近的操作须在遮挡不活跃的区间内执行。'
        '需要指出，遮挡仅在通信可用期间有实际影响；若遮挡区间落在通信中断期内，则为无效遮挡。')

    doc.add_heading('2.3.4 提前量约束', level=3)
    doc.add_paragraph('大臂和小臂的重启加电保温需满足操作规程要求的提前量：')
    fm(doc, 'end(τ_BA) + 43200 ≤ start(τ_P3)，end(τ_SA) + 18000 ≤ start(τ_P4)                (5)')

    p = doc.add_paragraph(); br(p, '表2-1 约束类型汇总')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    htb(doc, ['约束类型', '数学表达', '适用范围'],
        [['持续时间', 'end(τ)=start(τ)+dur(τ)', '所有子任务'],
         ['紧邻顺序', 'end(τᵢ)=start(τⱼ)', '阶段内子任务'],
         ['通信窗口', '∃k: aₖ≤start∧end≤bₖ', '所有操作子任务'],
         ['遮挡回避', '∃k: σₖ=Inactive∧Iₖ⊇[s,e]', 'SWA附近子任务'],
         ['大臂提前', 'end(τ_BA)+43200≤start(τ_P3)', '大臂重启'],
         ['小臂提前', 'end(τ_SA)+18000≤start(τ_P4)', '小臂重启']])
    doc.add_paragraph()

    # ═══════════════════
    # 2.4
    # ═══════════════════
    doc.add_heading('2.4 NDDL建模编码', level=2)

    doc.add_paragraph(
        '基于上述形式化模型，使用EUROPA框架的NDDL语言进行编码实现。'
        'NDDL是一种声明式建模语言，以类继承和谓词定义为基础，支持时间线、动作规则和约束的自然表达。'
        '模型定义三个Timeline子类，分别对应定义1中的三个对象：'
        'CommWindow定义通信可用和中断两种状态；OcclusionSWA定义遮挡活跃和不活跃两种状态；'
        'Arm定义机械臂的各阶段动作Token和完成标志。')

    doc.add_paragraph(
        '其中平台状态设置阶段（Phase 2）因总时长3530秒超过单个通信窗口容量（约3400秒），'
        '需拆分为Phase2A（太阳帆板+舱外相机设置，2410秒）和Phase2B（禁止项设置，1120秒），'
        '分别在不同通信窗口内执行。这一拆分在NDDL中通过定义两个独立的action Token实现。')

    # ═══════════════════
    # 2.5
    # ═══════════════════
    doc.add_heading('2.5 动作规则编码', level=2)

    doc.add_paragraph(
        '每个action的规则体定义了该动作被激活时自动创建的子Token和约束，'
        '对应定义1中的规则集合ℛ。以Phase 4（小臂运动至SWA）为例，'
        '该阶段同时涉及持续时间、顺序、通信和遮挡四类约束，是约束最完整的典型动作。'
        '其规则体的含义为：Phase 4的持续时间固定为770秒；Phase 3必须紧邻先于Phase 4执行；'
        'Phase 4的执行时段须完全包含在某个通信可用窗口内；'
        '同时须完全包含在某个遮挡不活跃的区间内。')

    doc.add_paragraph(
        '当求解器激活Phase 4 Token时，规则引擎自动执行以下操作：'
        '（1）在STN中添加持续时间约束的上界边和下界边；'
        '（2）创建Phase 3子Token（若不存在则为新的开放条件缺陷），添加紧邻约束；'
        '（3）创建InComms子Token，添加通信窗口包含约束；'
        '（4）创建Inactive子Token，添加遮挡回避包含约束。'
        '步骤（2）（3）（4）创建的子Token均以未激活状态进入计划，成为新的缺陷，驱动求解器继续推理。')

    # ═══════════════════
    # 2.6
    # ═══════════════════
    doc.add_heading('2.6 求解算法', level=2)

    doc.add_paragraph(
        '基于2.2节建立的形式化模型和2.3节定义的约束体系，本节给出求解算法的完整流程。'
        '求解算法的核心是缺陷导向的时序回溯搜索，分为三个阶段：约束网络构建、缺陷导向搜索和计划提取。'
        '图2-1给出了完整的求解流程。')

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(); r.add_picture(f'{OUT}/europa_flowchart.png', width=Inches(4.5))
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.add_run('图2-1 EUROPA求解机械臂任务规划流程').italic = True
    doc.add_paragraph()

    # ── 2.6.1 ──
    doc.add_heading('2.6.1 约束网络构建', level=3)

    doc.add_paragraph(
        '约束网络构建对应图2-1中Phase 1和Phase 2，其步骤如下：')

    p = doc.add_paragraph(); br(p, '算法1. 约束网络构建算法（BUILD_STN）')

    doc.add_paragraph(
        '1）输入NDDL域模型M和问题实例P，通过ANTLR3解析器将NDDL文件转换为内部表示，'
        '注册所有类型定义、谓词和规则到Schema中；')
    doc.add_paragraph(
        '2）根据问题实例创建对象和初始Token，包括通信窗口序列、遮挡状态序列'
        '以及Phase 1子任务的固定放置，构成初始部分计划PP₀；')
    doc.add_paragraph(
        '3）初始化STN N=(V, E, O)，为PP₀中每个Token的start、end、duration变量创建时间点节点，'
        '按公式(2)将持续时间约束编码为上界边和下界边；')
    doc.add_paragraph(
        '4）将Fact Token的固定时间值作为约束添加到STN中，即添加原点O到对应时间点的距离约束；')
    doc.add_paragraph(
        '5）执行Bellman-Ford算法计算所有节点的势函数π(v)，同时检测是否存在负权环路；'
        '若存在负环，说明初始约束已矛盾，返回失败；否则输出初始STN N和部分计划PP₀。')

    # ── 2.6.2 ──
    doc.add_heading('2.6.2 缺陷导向搜索主循环', level=3)

    doc.add_paragraph(
        '缺陷导向搜索是求解算法的核心，对应图2-1中Phase 3至Phase 7。'
        '该过程维护一个决策栈记录已做出的决策，在需要回溯时支持逐层撤销。')

    p = doc.add_paragraph(); br(p, '算法2. 缺陷导向时序回溯搜索算法（Solve）')

    doc.add_paragraph(
        '1）输入初始STN N、初始部分计划PP₀和最大搜索步数N_max，初始化决策栈为空，步计数为零；')
    doc.add_paragraph(
        '2）执行STN约束传播（算法5），验证当前约束系统的一致性；若初始状态即不一致，返回失败；')
    doc.add_paragraph(
        '3）在当前部分计划中选择一个缺陷进行修复（算法3）；若无缺陷可选，说明所有缺陷已消除，'
        '当前部分计划即为完整计划，返回成功；')
    doc.add_paragraph(
        '4）根据所选缺陷的类型，生成决策点并枚举可选方案：'
        '对开放条件缺陷，可选方案包括激活、合并和拒绝；'
        '对时序威胁缺陷，可选方案为时间线上的合法排序位置；'
        '对未绑定变量缺陷，可选方案为变量域中的候选值；')
    doc.add_paragraph(
        '5）执行所选方案，修改部分计划和STN；'
        '若决策类型为激活（ACTIVATE），则触发规则引擎——'
        '遍历该Token类型关联的所有规则，自动创建子Token和约束（2.5节），'
        '这些新Token成为新的缺陷，驱动搜索继续向前推进；')
    doc.add_paragraph(
        '6）再次执行STN约束传播，验证决策后的一致性；'
        '若一致，将决策压入栈中，转步骤2继续下一轮迭代；')
    doc.add_paragraph(
        '7）若传播后不一致，执行时序回溯（算法4）：从当前失败的决策开始，'
        '逐层撤销已做出的决策并尝试替代方案，直到找到可继续的分支或搜索空间耗尽；')
    doc.add_paragraph(
        '8）重复步骤2至步骤7，直到求解成功、搜索空间耗尽或达到最大步数限制。')

    # ── 2.6.3 ──
    doc.add_heading('2.6.3 缺陷选择策略', level=3)

    p = doc.add_paragraph(); br(p, '算法3. 缺陷选择算法（Allocate_Decision）')

    doc.add_paragraph(
        '1）输入当前部分计划中的缺陷管理器集合（按优先级排列为'
        'ThreatManager、OpenConditionManager和UnboundVariableManager）；')
    doc.add_paragraph(
        '2）首先在所有管理器中查找零承诺缺陷——即仅有唯一解决方案的缺陷'
        '（例如规则触发创建的子Token只能合并到特定的Fact Token）；'
        '若找到，直接返回该缺陷的决策点，此类决策不引入搜索分支，可无代价地推进求解过程；')
    doc.add_paragraph(
        '3）若无零承诺缺陷，则遍历所有管理器中的候选缺陷，'
        '根据管理器配置优先级和缺陷自身的紧急度综合评分，选择优先级最高的缺陷，'
        '创建并初始化其决策点后返回。')

    # ── 2.6.4 ──
    doc.add_heading('2.6.4 时序回溯', level=3)

    p = doc.add_paragraph(); br(p, '算法4. 时序回溯算法（Backtrack）')

    doc.add_paragraph(
        '1）从当前失败的决策或决策栈顶取出一个决策点；')
    doc.add_paragraph(
        '2）若该决策已执行，则撤销其对部分计划和STN的修改——'
        '移除该决策添加的约束边，恢复被修改Token的状态；')
    doc.add_paragraph(
        '3）检查该决策是否还有未尝试的替代方案；'
        '若有，保留该决策点作为当前活跃决策，回溯终止，搜索从替代方案继续；')
    doc.add_paragraph(
        '4）若该决策的所有方案已耗尽，删除该决策点，继续从栈中弹出上一层决策，重复步骤2；')
    doc.add_paragraph(
        '5）若决策栈已清空且无可用方案，返回搜索空间耗尽标志。')

    doc.add_paragraph(
        '回溯的关键代价在于约束删除后的重传播：由于移除约束可能使之前被排除的路径重新可行，'
        'STN需执行完全重传播（Bellman-Ford重算势函数 + 双向Dijkstra更新时间界），'
        '其复杂度为O(VE+(V+E)log V)。')

    # ── 2.6.5 ──
    doc.add_heading('2.6.5 STN增量传播', level=3)

    doc.add_paragraph(
        'STN增量传播是保证求解效率的关键机制。由于下界边的权重为负值（公式2），'
        '标准Dijkstra算法不可直接使用。本算法采用Johnson重标号技术：'
        '首先通过Bellman-Ford算法计算每个节点的势函数π(v)，'
        '然后将边权重标号为w\'(u,v) = w(u,v) + π(u) − π(v) ≥ 0，'
        '使得所有边权非负后即可使用Dijkstra算法进行高效传播。')

    p = doc.add_paragraph(); br(p, '算法5. STN增量传播算法（INC_PROPAGATE）')

    doc.add_paragraph(
        '1）输入STN N以及新约束影响的两个端点src和targ；')
    doc.add_paragraph(
        '2）阶段一（一致性检查）：通过增量Bellman-Ford检测新约束是否引入负权环路；'
        '若检测到负环，标记约束系统不一致，返回失败；')
    doc.add_paragraph(
        '3）阶段二（上界传播）：执行前向Dijkstra——'
        '从受影响的起始节点出发，沿出边方向传播，'
        '若发现新距离小于节点当前上界，则更新上界并将该节点以Johnson重标号后的非负键值插入优先队列，'
        '持续处理直至队列为空（算法6）；')
    doc.add_paragraph(
        '4）阶段三（下界传播）：执行后向Dijkstra——'
        '从受影响的节点出发，沿入边反向传播，更新受影响节点的下界'
        '（下界以负值编码，传播(−lowerBound)沿入边方向的最短路径），'
        '持续处理直至队列为空（算法7）；')
    doc.add_paragraph(
        '5）三个阶段执行前均通过START_NODE函数进行可行性预检：'
        '若新约束不能改进任一端点的距离值，则该阶段在O(1)内终止，'
        '避免无效的全图遍历。这一短路机制是增量传播在绝大多数情况下远快于完全重传播的根本原因。')

    p = doc.add_paragraph(); br(p, '算法6. STN前向传播算法（Dijkstra_Forward）')

    doc.add_paragraph(
        '1）输入含初始受影响节点的优先队列Queue；')
    doc.add_paragraph(
        '2）弹出队列中键值最小的节点u，遍历其所有出边(u,v,w)，'
        '计算候选距离newDist = u.upperBound + w；')
    doc.add_paragraph(
        '3）若newDist小于v的当前上界，则更新v.upperBound ← newDist，'
        '更新搜索深度depth(v) ← depth(u)+1，'
        '以Johnson重标号后的键值key = newDist − π(v)将v插入队列；')
    doc.add_paragraph(
        '4）若depth(v)超过节点总数|V|，则隐含存在负环，标记不一致并终止；')
    doc.add_paragraph(
        '5）重复步骤2至步骤4，直至队列为空，所有受影响节点的上界更新完毕。')

    p = doc.add_paragraph(); br(p, '算法7. STN后向传播算法（Dijkstra_Backward）')

    doc.add_paragraph(
        '1）输入含初始受影响节点的优先队列Queue；')
    doc.add_paragraph(
        '2）弹出队列中键值最小的节点u，遍历其所有入边(v,u,w)——注意此处沿入边反向遍历；')
    doc.add_paragraph(
        '3）计算候选距离newDist = (−u.lowerBound) + w，'
        '若newDist < (−v.lowerBound)，则更新v.lowerBound ← −newDist，'
        '以键值key = newDist + π(v)将v插入队列；')
    doc.add_paragraph(
        '4）重复直至队列为空，所有受影响节点的下界更新完毕。')

    # ── 2.6.6 ──
    doc.add_heading('2.6.6 计划提取与验证', level=3)

    p = doc.add_paragraph(); br(p, '算法8. 计划提取算法（EXTRACT_PLAN）')

    doc.add_paragraph(
        '1）输入求解完成的完整计划PP*和STN N；')
    doc.add_paragraph(
        '2）按时间线上的Token顺序遍历所有已激活Token，'
        '提取每个Token的时间下界（lowerBound）作为最早开始时间，'
        '加上固定持续时间得到结束时间；')
    doc.add_paragraph(
        '3）验证每个操作Token的时间区间[start, end]是否被包含在某个通信可见窗口[aₖ, bₖ]内（公式4），'
        '以及提前量约束（公式5）是否满足；')
    doc.add_paragraph(
        '4）输出可执行的任务计划，包含各子任务的起止时间、所属通信窗口和约束满足状态。')

    # ── 2.6.7 ──
    doc.add_heading('2.6.7 复杂度分析', level=3)

    p = doc.add_paragraph(); br(p, '表2-2 算法各阶段时间复杂度')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    htb(doc, ['阶段', '核心算法', '时间复杂度', '说明'],
        [['约束网络构建', 'Bellman-Ford', 'O(VE)', 'V为时间点数，E为边数'],
         ['增量传播(添加)', 'Johnson+Dijkstra', 'O((V+E)log V)', '每次决策后执行'],
         ['增量传播(删除)', 'BF+Dijkstra', 'O(VE+(V+E)log V)', '回溯时执行'],
         ['缺陷选择', '线性扫描', 'O(|Flaws|)', '缺陷遍历与评分'],
         ['完整求解', 'S次迭代', 'O(S·(V+E)log V)', 'S为搜索步数']])
    doc.add_paragraph()

    doc.add_paragraph(
        '对于本文的机械臂算例，V≈80，E≈160，S≈200，完整求解在亚秒级时间内完成。'
        '与显式状态空间搜索方法（如UPMurphi的前向搜索）相比，'
        'EUROPA的缺陷导向搜索直接在部分计划空间中操作，'
        '通过STN增量传播实现高效的约束推理，避免了显式枚举可达状态空间的指数级开销。')

    # ═══════════════════
    # 2.7
    # ═══════════════════
    doc.add_heading('2.7 本章小结', level=2)

    doc.add_paragraph(
        '本章对空间站机械臂在轨作业任务规划问题进行了系统的形式化建模与求解算法设计。'
        '2.2节以五个形式化定义建立了时间线规划系统的数学框架，'
        '各定义按TPS→部分计划→缺陷→STN→规划问题的逻辑递进排列。'
        '2.3节对六类约束进行了统一的数学建模（公式3-5）。'
        '2.4至2.5节使用NDDL语言完成了模型编码和动作规则设计。'
        '2.6节以步骤化说明形式给出了八个关键算法的完整描述：'
        '算法1构建约束网络，算法2实现搜索主循环，算法3-4处理缺陷选择和回溯，'
        '算法5-7实现STN的三阶段增量传播（一致性检查、上界传播、下界传播），'
        '算法8完成计划提取与验证。'
        'STN增量Dijkstra传播将单次约束推理控制在O((V+E)log V)，'
        '使得包含约200步搜索的完整求解在亚秒级时间内完成，满足空间站机械臂任务规划的实时性要求。')

    # Save
    for path in [f'{OUT}/chapter2_v2.docx', '/workspace/documentation/chapter2_v2.docx']:
        doc.save(path)
    print(f'Saved: {path}')

if __name__ == '__main__':
    create()
